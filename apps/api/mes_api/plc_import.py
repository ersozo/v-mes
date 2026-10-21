"""TIA Portal DB tablosundan (Excel dışa aktarımı veya DB editöründen kopyala-yapıştır) tag adayları çıkarır.

Sütun düzeni dosyadan dosyaya değiştiği için sütunlar başlık adına değil İÇERİĞE göre tanınır:
    ofset   : "892.1" biçimi (optimize edilmemiş DB'de TIA'nın gösterdiği bayt.bit)
    tip     : TIA veri tipi (Bool, Int, UInt, DInt, Real, Struct, Array[..] of ..., String ...)
    ad      : satırdaki diğer ilk metin
Struct satırları (ör. "MakinaIstasyonData[2]") ardından gelen alanlara ad öneki olur.
Tipi boş bırakılmış ve bit ofseti 0'dan farklı satırlar Bool kabul edilir (ve işaretlenir).

Çıktı yalnızca ÖNERİDİR: kullanıcı satırları seçer, ad ve rol atar; asıl doğrulama mes_api.plc ile yapılır.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import asdict, dataclass

MAX_ROWS = 5000
MAX_BYTES = 2 * 1024 * 1024

# TIA tipi → node-red-contrib-s7 tipi (None = desteklenmiyor, nedeniyle)
_TYPE_MAP: dict[str, str | None] = {
    "bool": "X",
    "byte": "BYTE", "usint": "BYTE",
    "int": "INT",
    "word": "WORD", "uint": "WORD",
    "dint": "DINT",
    "dword": "DWORD", "udint": "DWORD",
    "real": "REAL",
}
_UNSUPPORTED: dict[str, str] = {
    "sint": "8 bit işaretli tamsayı desteklenmiyor",
    "lint": "64 bit tamsayı desteklenmiyor", "ulint": "64 bit tamsayı desteklenmiyor", "lword": "64 bit desteklenmiyor",
    "lreal": "LReal (64 bit) desteklenmiyor",
    "char": "metin desteklenmiyor", "wchar": "metin desteklenmiyor",
    "time": "Time desteklenmiyor", "date": "Date desteklenmiyor", "tod": "Time_Of_Day desteklenmiyor",
    "time_of_day": "Time_Of_Day desteklenmiyor", "dtl": "DTL desteklenmiyor", "date_and_time": "Date_And_Time desteklenmiyor",
    "s5time": "S5Time desteklenmiyor",
}
_OFFSET_RE = re.compile(r"^(\d{1,5})\.([0-7])$")
_DB_RE = re.compile(r"^%?db\s*(\d{1,5})(?:\b|\.|$)", re.IGNORECASE)
_ARRAY_RE = re.compile(r"^array\s*\[.*\]\s*of\s+(.+)$", re.IGNORECASE)
_STRING_RE = re.compile(r"^w?string(\s*\[\d+\])?$", re.IGNORECASE)
_HEADER_WORDS = {"static", "input", "output", "inout", "temp", "constant", "name", "ad", "data type", "datatype", "offset", "ofset", "type"}


@dataclass
class ImportRow:
    line: int                  # kaynak satır numarası (1'den)
    path: str                  # TIA yolu, ör. MakinaIstasyonData[2].girisSensoru
    name: str                  # önerilen tag adı
    tia_type: str | None
    offset: str
    address: str | None        # ör. DB200,X902.0 (desteklenmiyorsa None)
    supported: bool
    note: str | None = None
    suggested_role: str = "process"
    suggested_station: int | None = None   # dizi elemanındaki koşu biti → istasyon numarası


def _translit(text: str) -> str:
    text = text.replace("ı", "i").replace("İ", "I")
    return "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))


def suggest_name(path: str) -> str:
    """'MakinaIstasyonData[2].girisSensoru' → 'makina_istasyon_data_2_giris_sensoru' (tag adı kuralına uygun)."""
    s = _translit(path)
    s = re.sub(r"\[(\d+)\]", r"_\1", s)
    s = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", s)
    s = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1_\2", s)
    s = re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_").lower()
    if not s or not s[0].isalpha():
        s = "t_" + s
    return s[:64].rstrip("_")


def _suggest_role(name: str, s7_type: str | None) -> str:
    n = name.lower()
    # Hurda önce: "HurdaAdet" hem "hurda" hem "adet" içerir
    if s7_type in ("WORD", "DWORD", "DINT") and re.search(r"hurda|fire|scrap|reject|(^|[_.])nok([_.]|$)|red_?adet|ret_?adet", n):
        return "scrap_count_candidate"
    if s7_type in ("WORD", "DWORD", "DINT") and re.search(r"sayac|sayici|counter|count|adet|toplam", n):
        return "total_count_candidate"
    if s7_type == "X" and re.search(r"kosu|calis|run|auto|otomatik", n):
        return "running_candidate"
    return "process"


def _suggest_station(path: str, role: str) -> int | None:
    """'MakinaIstasyonData[3].kosuBiti' → 3 (yalnızca koşu/çalışıyor biti adaylarında)."""
    if role != "running_candidate":
        return None
    m = re.search(r"\[(\d+)\]\.[^.\[\]]+$", path)
    return int(m.group(1)) if m and 1 <= int(m.group(1)) <= 99 else None


def _classify(cells: list[str]) -> tuple[str | None, str | None, str | None]:
    """(ad, tip, ofset) — içerik bazlı."""
    offset = next((c for c in cells if _OFFSET_RE.match(c)), None)
    type_cell = None
    for c in cells:
        lc = c.lower()
        if lc in _TYPE_MAP or lc in _UNSUPPORTED or lc == "struct" or _ARRAY_RE.match(c) or _STRING_RE.match(c):
            type_cell = c
            break
    name = next((c for c in cells if c not in (offset, type_cell) and not _DB_RE.match(c) and c.lower() not in _HEADER_WORDS
                 and not re.match(r"^[\d.,\-+eE]+$", c) and c.lower() not in ("true", "false")), None)
    return name, type_cell, offset


def parse_rows(rows: list[list[object]], db_number: int | None = None) -> dict:
    warnings: list[str] = []
    detected_db = None
    out: list[ImportRow] = []
    prefix: str | None = None

    for i, raw in enumerate(rows[:MAX_ROWS], start=1):
        cells = [str(c).strip() for c in raw if c is not None and str(c).strip() != ""]
        if not cells:
            continue
        if detected_db is None:
            m = next((_DB_RE.match(c) for c in cells if _DB_RE.match(c)), None)
            if m:
                detected_db = int(m.group(1))
        name, tia_type, offset = _classify(cells)
        if not name or not offset:
            continue   # başlık / boş / açıklama satırı
        lt = (tia_type or "").lower()

        if lt == "struct":
            prefix = name
            continue
        if _ARRAY_RE.match(tia_type or ""):
            elem = _ARRAY_RE.match(tia_type or "").group(1).strip().lower()  # type: ignore[union-attr]
            if elem == "struct":
                continue   # elemanları ayrı Struct satırları olarak gelir
            out.append(ImportRow(i, name, suggest_name(name), tia_type, offset, None, False,
                                 "Dizi: elemanları ayrı satırlar olarak dışa aktarın"))
            continue

        path = f"{prefix}.{name}" if prefix and not name.startswith(prefix) else name
        byte, bit = (int(x) for x in _OFFSET_RE.match(offset).groups())  # type: ignore[union-attr]
        note = None
        if not tia_type:
            if bit != 0:
                tia_type, lt, note = "Bool", "bool", "Tip boş; bit ofsetinden Bool varsayıldı — doğrulayın"
            else:
                out.append(ImportRow(i, path, suggest_name(path), None, offset, None, False, "Tip boş; TIA'dan veri tipini kontrol edin"))
                continue

        if _STRING_RE.match(tia_type):
            out.append(ImportRow(i, path, suggest_name(path), tia_type, offset, None, False, "Metin (String) şimdilik okunmuyor"))
            continue
        s7 = _TYPE_MAP.get(lt)
        if s7 is None:
            out.append(ImportRow(i, path, suggest_name(path), tia_type, offset, None, False, _UNSUPPORTED.get(lt, f"Bilinmeyen tip: {tia_type}")))
            continue
        if s7 != "X" and bit != 0:
            out.append(ImportRow(i, path, suggest_name(path), tia_type, offset, None, False, "Bool dışındaki tipte bit ofseti 0 olmalı"))
            continue

        db = db_number or detected_db
        address = None if db is None else (f"DB{db},X{byte}.{bit}" if s7 == "X" else f"DB{db},{s7}{byte}")
        role = _suggest_role(path, s7)
        out.append(ImportRow(i, path, suggest_name(path), tia_type, offset, address, address is not None,
                             note or (None if db else "DB numarası bulunamadı — içe aktarma ekranında girin"),
                             role, _suggest_station(path, role)))

    if len(rows) > MAX_ROWS:
        warnings.append(f"İlk {MAX_ROWS} satır işlendi")
    if not (db_number or detected_db):
        warnings.append("DB numarası dosyada bulunamadı; elle girin")
    names = [r.name for r in out]
    for dup in sorted({n for n in names if names.count(n) > 1}):
        warnings.append(f'"{dup}" adı birden fazla satırda öneriliyor; içe aktarmadan önce yeniden adlandırın')
    if not out:
        warnings.append("Ofseti olan satır bulunamadı. DB optimize edilmiş olabilir (ofset sütunu yok) veya biçim tanınmadı.")
    return {"db_number": db_number or detected_db, "rows": [asdict(r) for r in out], "warnings": warnings}


def read_xlsx(data: bytes) -> list[list[object]]:
    from openpyxl import load_workbook   # defusedxml kuruluysa openpyxl onu kullanır (XML saldırılarına karşı)

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        rows: list[list[object]] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) > MAX_ROWS:
                    return rows
        return rows
    finally:
        wb.close()


def read_text(text: str) -> list[list[object]]:
    """TIA DB editöründen kopyalanan (sekmeyle ayrılmış) veya CSV (; veya ,) metin."""
    sample = text[:2000]
    delim = "\t" if "\t" in sample else ";" if sample.count(";") >= sample.count(",") else ","
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)][: MAX_ROWS + 1]
